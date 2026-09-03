from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
W = 16.2

STYLES = {
    "clean_blue": {"name": "A｜专业蓝白·证据型", "p": "2F80ED", "soft": "EAF3FF", "paper": "F7FAFE", "text": "243447", "muted": "667085", "line": "D9E5F2"},
    "teal_grid": {"name": "B｜现代青绿·双栏型", "p": "147D7A", "soft": "EAF7F6", "paper": "F6FBFB", "text": "20363A", "muted": "66797C", "line": "D7E9E8"},
    "indigo_report": {"name": "C｜稳重靛蓝·报告型", "p": "4B4E9B", "soft": "F0F0FB", "paper": "F8F8FC", "text": "292B3D", "muted": "6D7085", "line": "DDDDF0"},
    "amber_standard": {"name": "D｜商务金棕·杂志型", "p": "A87520", "soft": "FBF4E7", "paper": "FCFAF5", "text": "342D24", "muted": "776F63", "line": "E9DFC9"},
    "green_governance": {"name": "E｜治理绿白·卡片型", "p": "3C7D3B", "soft": "EEF7ED", "paper": "F8FCF7", "text": "263529", "muted": "6B786C", "line": "D9E8D7"},
}


def theme(style):
    if style not in STYLES:
        raise ValueError(style)
    return STYLES[style]


def text(p, s, size=10, bold=False, color="243447"):
    r = p.add_run(s)
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    r.bold = bold
    return r


def fill(cell, color):
    pr = cell._tc.get_or_add_tcPr(); node = pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd"); pr.append(node)
    node.set(qn("w:fill"), color)


def line(cell, color, size="6"):
    pr = cell._tc.get_or_add_tcPr(); bs = pr.first_child_found_in("w:tcBorders")
    if bs is None:
        bs = OxmlElement("w:tcBorders"); pr.append(bs)
    for e in ("top", "left", "bottom", "right"):
        n = bs.find(qn("w:" + e))
        if n is None:
            n = OxmlElement("w:" + e); bs.append(n)
        n.set(qn("w:val"), "single"); n.set(qn("w:sz"), size); n.set(qn("w:color"), color)


def pad(cell, v=120, h=120):
    pr = cell._tc.get_or_add_tcPr(); mar = pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar"); pr.append(mar)
    for k, x in (("top", v), ("bottom", v), ("start", h), ("end", h)):
        n = mar.find(qn("w:" + k))
        if n is None:
            n = OxmlElement("w:" + k); mar.append(n)
        n.set(qn("w:w"), str(x)); n.set(qn("w:type"), "dxa")


def set_widths(table, values):
    table.autofit = False
    for i, v in enumerate(values):
        if i < len(table.columns):
            table.columns[i].width = Cm(v)


def image(cell, path, t, width=None, height=None, label="产品界面截图"):
    fill(cell, t["paper"]); line(cell, t["line"]); pad(cell, 120, 100)
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if path and Path(path).exists():
        if width: p.add_run().add_picture(path, width=Cm(width))
        else: p.add_run().add_picture(path, height=Cm(height or 8))
    else:
        text(p, label, 9, False, t["muted"])


def heading(doc, page, t, style, compact=False):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1)
    text(p, page.get("code", ""), 8, True, t["p"])
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
    text(p, page.get("title", ""), 20 if not compact else 18, True, t["text"])
    if page.get("summary"):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(7)
        text(p, page["summary"], 9.7, False, t["muted"])


def section(doc, label, t, prefix=""):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(4)
    text(p, prefix + label, 11.2, True, t["p"])


def value_bar(doc, values, t, mode="strip"):
    if not values: return
    table = doc.add_table(rows=1, cols=len(values)); set_widths(table, [W/len(values)]*len(values))
    for i, v in enumerate(values):
        c = table.cell(0, i); fill(c, t["soft"] if mode != "outline" else "FFFFFF"); line(c, t["line"]); pad(c, 90, 90)
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        text(p, ("● " if mode == "strip" else f"0{i+1}  ") + v, 8.5, True, t["p"])


def feature_table(doc, feats, t, compact=False):
    table = doc.add_table(rows=1, cols=2); set_widths(table, [4.1, 12.1])
    for i, h in enumerate(("功能", "说明")):
        c=table.rows[0].cells[i]; fill(c,t["p"]); pad(c,70,100); text(c.paragraphs[0],h,8.8,True,"FFFFFF")
    for f in feats:
        cs=table.add_row().cells
        for c in cs: line(c,t["line"]); pad(c,65,95); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        fill(cs[0],t["soft"]); text(cs[0].paragraphs[0],f.get("name",""),9,True,t["p"])
        text(cs[1].paragraphs[0],f.get("description",""),8.6 if compact else 9,t["text"])
    return table


def cards(doc, feats, t, cols=3, numbered=False):
    if not feats: return
    rows=(len(feats)+cols-1)//cols; table=doc.add_table(rows=rows,cols=cols); set_widths(table,[W/cols]*cols)
    for i,f in enumerate(feats):
        c=table.cell(i//cols,i%cols); fill(c,t["paper"]); line(c,t["line"]); pad(c,125,120)
        p=c.paragraphs[0]; text(p,(f"{i+1:02d}  " if numbered else "")+f.get("name","能力"),9.3,True,t["p"])
        p=c.add_paragraph(); text(p,f.get("description",""),8.5,False,t["text"])


def layout_standard(doc,page,t,style):
    feats=page.get("features",[])
    if style=="clean_blue":
        tb=doc.add_table(rows=1,cols=1); set_widths(tb,[W]); image(tb.cell(0,0),page.get("screenshot"),t,width=15.0)
        section(doc,"功能说明",t); feature_table(doc,feats,t,True); section(doc,"应用价值",t); value_bar(doc,page.get("values",[]),t)
    elif style=="teal_grid":
        tb=doc.add_table(rows=1,cols=2); set_widths(tb,[10.3,5.9]); image(tb.cell(0,0),page.get("screenshot"),t,width=9.6)
        c=tb.cell(0,1); fill(c,t["soft"]); line(c,t["line"]); pad(c,150,125); text(c.paragraphs[0],"核心能力",11,True,t["p"])
        for i,f in enumerate(feats[:5],1):
            p=c.add_paragraph(); text(p,f"0{i}  {f.get('name','')}",9.1,True,t["text"])
            p=c.add_paragraph(); text(p,f.get("description",""),8.2,False,t["muted"])
        section(doc,"应用价值",t,"▌ "); value_bar(doc,page.get("values",[]),t)
    elif style=="indigo_report":
        cards(doc,feats[:3],t,3,True); section(doc,"产品界面证据",t,"01  ")
        tb=doc.add_table(rows=1,cols=1); set_widths(tb,[W]); image(tb.cell(0,0),page.get("screenshot"),t,width=14.5)
        if len(feats)>3: section(doc,"详细功能",t); feature_table(doc,feats[3:],t,True)
        section(doc,"管理价值",t,"02  "); value_bar(doc,page.get("values",[]),t,"outline")
    elif style=="amber_standard":
        tb=doc.add_table(rows=1,cols=2); set_widths(tb,[10.8,5.4]); image(tb.cell(0,0),page.get("screenshot"),t,width=10.0)
        c=tb.cell(0,1); fill(c,"FFFFFF"); line(c,t["line"]); pad(c,100,110)
        for i,f in enumerate(feats[:5],1):
            p=c.paragraphs[0] if i==1 else c.add_paragraph(); text(p,f"{i:02d}",15,True,t["p"]); text(p,"  "+f.get("name",""),9.1,True,t["text"])
            p=c.add_paragraph(); text(p,f.get("description",""),8.2,False,t["muted"])
        section(doc,"应用价值",t,"01  "); value_bar(doc,page.get("values",[]),t,"outline")
    else:
        cards(doc,feats[:6],t,3,True); section(doc,"界面证据",t)
        tb=doc.add_table(rows=1,cols=1); set_widths(tb,[W]); image(tb.cell(0,0),page.get("screenshot"),t,width=14.4)
        section(doc,"管理价值",t); value_bar(doc,page.get("values",[]),t)


def layout_annotated(doc,page,t,style):
    heading(doc,page,t,style,True); feats=page.get("features",[])[:6]
    tb=doc.add_table(rows=1,cols=1); set_widths(tb,[W]); image(tb.cell(0,0),page.get("screenshot"),t,width=14.7,label="PC 截图 + ①②③功能标注")
    section(doc,"截图标注说明",t)
    cols=2 if style in ("teal_grid","green_governance") else 1
    cards(doc,feats,t,cols,True)
    value_bar(doc,page.get("values",[]),t)


def layout_pc_mobile(doc,page,t,style):
    heading(doc,page,t,style,True); pc=page.get("pc_screenshot") or page.get("screenshot"); mob=page.get("mobile_screenshot")
    if style=="clean_blue":
        tb=doc.add_table(rows=1,cols=2); set_widths(tb,[11.2,5.0]); image(tb.cell(0,0),pc,t,width=10.3,label="管理端"); image(tb.cell(0,1),mob,t,height=8.7,label="员工端")
    elif style=="teal_grid":
        tb=doc.add_table(rows=1,cols=2); set_widths(tb,[8.1,8.1]); image(tb.cell(0,0),pc,t,width=7.2,label="管理端"); image(tb.cell(0,1),mob,t,height=8.7,label="员工端")
    elif style=="indigo_report":
        tb=doc.add_table(rows=2,cols=2); set_widths(tb,[8.1,8.1]); image(tb.cell(0,0),pc,t,width=7.4,label="管理端"); image(tb.cell(0,1),mob,t,height=8.2,label="员工端")
        c=tb.cell(1,0); fill(c,t["soft"]); line(c,t["line"]); pad(c,100,100); text(c.paragraphs[0],"管理端",9,True,t["p"])
        c=tb.cell(1,1); fill(c,t["soft"]); line(c,t["line"]); pad(c,100,100); text(c.paragraphs[0],"员工端",9,True,t["p"])
    elif style=="amber_standard":
        tb=doc.add_table(rows=1,cols=2); set_widths(tb,[12.0,4.2]); image(tb.cell(0,0),pc,t,width=11.1,label="管理端"); image(tb.cell(0,1),mob,t,height=8.8,label="员工端")
    else:
        section(doc,"双端协同",t); cards(doc,[{"name":"管理端","description":"配置、审核、管理与数据查看"},{"name":"员工端","description":"申请、预订、行程与移动服务"}],t,2,True)
        tb=doc.add_table(rows=1,cols=2); set_widths(tb,[10.5,5.7]); image(tb.cell(0,0),pc,t,width=9.7,label="管理端"); image(tb.cell(0,1),mob,t,height=8.5,label="员工端")
    section(doc,"功能说明",t); feature_table(doc,page.get("features",[])[:5],t,True); value_bar(doc,page.get("values",[]),t)


def layout_mobile(doc,page,t,style):
    heading(doc,page,t,style,True); shots=page.get("screenshots",[])[:3]; labels=page.get("mobile_labels",[])
    if style=="amber_standard":
        for i,path in enumerate(shots):
            tb=doc.add_table(rows=1,cols=2); set_widths(tb,[5.1,11.1]); image(tb.cell(0,0),path,t,height=8.5,label=labels[i] if i<len(labels) else "移动端"); c=tb.cell(0,1); fill(c,"FFFFFF"); pad(c,120,140); text(c.paragraphs[0],labels[i] if i<len(labels) else "移动端",11,True,t["p"]); text(c.add_paragraph(),page.get("summary",""),8.7,False,t["muted"])
    else:
        cols=3 if style in ("teal_grid","green_governance") else 2; n=len(shots); rows=(n+cols-1)//cols
        tb=doc.add_table(rows=rows,cols=cols); set_widths(tb,[W/cols]*cols)
        for i,path in enumerate(shots):
            c=tb.cell(i//cols,i%cols); image(c,path,t,height=8.7,label="手机端截图"); p=c.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; text(p,labels[i] if i<len(labels) else "移动端",8.5,True,t["p"])
    section(doc,"移动端能力",t); cards(doc,page.get("features",[])[:6],t,2 if style=="indigo_report" else 3,True); value_bar(doc,page.get("values",[]),t)


def layout_process(doc,page,t,style):
    heading(doc,page,t,style); steps=page.get("steps",[])[:6]
    if style=="indigo_report":
        section(doc,"业务流程",t,"01  ");
        for i,s in enumerate(steps):
            tb=doc.add_table(rows=1,cols=2); set_widths(tb,[2.0,14.2]); c=tb.cell(0,0); fill(c,t["p"]); pad(c,100,80); p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; text(p,f"{i+1:02d}",15,True,"FFFFFF"); c=tb.cell(0,1); fill(c,t["paper"]); line(c,t["line"]); pad(c,100,120); text(c.paragraphs[0],s.get("title","步骤"),9.4,True,t["text"]); text(c.add_paragraph(),s.get("description",""),8.5,False,t["muted"])
    else:
        cols=3 if style in ("teal_grid","green_governance") else 5 if style=="clean_blue" else 4; n=min(cols,len(steps)); tb=doc.add_table(rows=1,cols=n); set_widths(tb,[W/n]*n)
        for i,s in enumerate(steps[:n]):
            c=tb.cell(0,i); fill(c,t["soft"] if i%2==0 else t["paper"]); line(c,t["line"]); pad(c,120,85); p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; text(p,f"{i+1:02d}",16,True,t["p"]); p=c.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; text(p,s.get("title","步骤"),9.1,True,t["text"]); p=c.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; text(p,s.get("description",""),8.1,False,t["muted"])
    if page.get("screenshot"):
        section(doc,"关键界面",t); tb=doc.add_table(rows=1,cols=1); set_widths(tb,[W]); image(tb.cell(0,0),page["screenshot"],t,width=14.3)
    value_bar(doc,page.get("values",[]),t)


def layout_dashboard(doc,page,t,style):
    heading(doc,page,t); metrics=page.get("metrics",[])[:4]
    if style=="teal_grid":
        if metrics: cards(doc,[{"name":m.get("name","指标"),"description":m.get("value","-")} for m in metrics],t,4,False)
        tb=doc.add_table(rows=1,cols=1); set_widths(tb,[W]); image(tb.cell(0,0),page.get("screenshot"),t,width=14.7,label="BI 数据看板")
    elif style=="indigo_report":
        tb=doc.add_table(rows=1,cols=1); set_widths(tb,[W]); image(tb.cell(0,0),page.get("screenshot"),t,width=14.7,label="BI 数据看板")
        if metrics: value_bar(doc,[f"{m.get('name','指标')}  {m.get('value','-')}" for m in metrics],t,"outline")
    elif style=="green_governance":
        section(doc,"管理指标",t); cards(doc,[{"name":m.get("name","指标"),"description":m.get("value","-")} for m in metrics],t,2,True); section(doc,"数据证据",t); tb=doc.add_table(rows=1,cols=1); set_widths(tb,[W]); image(tb.cell(0,0),page.get("screenshot"),t,width=14.2,label="治理数据看板")
    else:
        tb=doc.add_table(rows=1,cols=1); set_widths(tb,[W]); image(tb.cell(0,0),page.get("screenshot"),t,width=14.8,label="BI 数据看板")
        if metrics: value_bar(doc,[f"{m.get('value','-')}  {m.get('name','指标')}" for m in metrics],t)
    if page.get("features"): section(doc,"指标解释",t); feature_table(doc,page["features"][:5],t,True)


def layout_overview(doc,page,t,style):
    heading(doc,page,t); items=page.get("items",[])
    if style=="indigo_report":
        section(doc,"能力地图",t,"01  "); cards(doc,items,t,2,True)
    elif style=="amber_standard":
        for i,item in enumerate(items):
            tb=doc.add_table(rows=1,cols=2); set_widths(tb,[2.3,13.9]); c=tb.cell(0,0); fill(c,t["p"]); pad(c,110,80); p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; text(p,f"{i+1:02d}",15,True,"FFFFFF"); c=tb.cell(0,1); fill(c,"FFFFFF"); line(c,t["line"]); pad(c,110,120); text(c.paragraphs[0],item.get("title","能力模块"),9.7,True,t["text"]); text(c.add_paragraph(),item.get("description",""),8.6,False,t["muted"])
    else:
        cards(doc,items,t,3 if style in ("teal_grid","green_governance") else 2,True)


def layout_matrix(doc,page,t,style):
    heading(doc,page,t); rows=page.get("rows",[]); headers=page.get("headers",["能力","员工端","管理端","数据"])
    tb=doc.add_table(rows=1,cols=len(headers)); set_widths(tb,[W/len(headers)]*len(headers))
    for i,h in enumerate(headers):
        c=tb.rows[0].cells[i]; fill(c,t["p"]); pad(c,80,90); text(c.paragraphs[0],h,8.8,True,"FFFFFF")
    for r in rows:
        cs=tb.add_row().cells
        for i,v in enumerate(r):
            line(cs[i],t["line"]); pad(cs[i],75,90); fill(cs[i],t["soft"] if i==0 else t["paper"]); text(cs[i].paragraphs[0],str(v),8.5,i==0,t["p"] if i==0 else t["text"])


def configure(doc,t,style):
    sec=doc.sections[0]; sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.7); sec.left_margin=Cm(2); sec.right_margin=Cm(2); sec.header_distance=Cm(.7); sec.footer_distance=Cm(.7)
    normal=doc.styles["Normal"]; normal.font.name="Microsoft YaHei"; normal.font.size=Pt(9.8); normal.font.color.rgb=RGBColor.from_string(t["text"])
    p=sec.header.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; text(p,"在途商旅 · "+t["name"],7.3,False,t["muted"])
    p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; text(p,"在途商旅产品功能手册",7.3,False,t["muted"])


def cover(doc,data,t):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(105); text(p,data.get("title","产品功能手册"),27,True,t["text"])
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; text(p,data.get("subtitle",""),11.5,False,t["muted"])
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(16); text(p,t["name"],9,True,t["p"])
    doc.add_page_break()


def render(data_path, output, style):
    data=json.loads(Path(data_path).read_text(encoding="utf-8")); t=theme(style); doc=Document(); configure(doc,t,style); cover(doc,data,t)
    for si,sec in enumerate(data.get("sections",[])):
        p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(70); text(p,sec.get("code",""),12,True,t["p"]); p=doc.add_paragraph(); text(p,sec.get("title",""),23,True,t["text"]); p=doc.add_paragraph(); text(p,sec.get("summary",""),9.8,False,t["muted"]); doc.add_page_break()
        for pi,page in enumerate(sec.get("pages",[])):
            kind=page.get("type","standard")
            heading_needed=kind in ("standard",)
            if heading_needed: heading(doc,page,t,style)
            if kind=="standard": layout_standard(doc,page,t,style)
            elif kind in ("annotated","screenshot_annotated"): layout_annotated(doc,page,t,style)
            elif kind=="pc_mobile": layout_pc_mobile(doc,page,t,style)
            elif kind=="mobile_gallery": layout_mobile(doc,page,t,style)
            elif kind in ("process","steps"): layout_process(doc,page,t,style)
            elif kind=="dashboard": layout_dashboard(doc,page,t,style)
            elif kind=="overview": layout_overview(doc,page,t,style)
            elif kind in ("matrix","capability_matrix"): layout_matrix(doc,page,t,style)
            else: layout_standard(doc,page,t,style)
            if not (si==len(data.get("sections",[]))-1 and pi==len(sec.get("pages",[]))-1): doc.add_page_break()
    output=Path(output); output.parent.mkdir(parents=True,exist_ok=True); doc.save(output); print(f"Generated: {output}")


def main():
    ap=argparse.ArgumentParser(description="在途商旅 Word Master Template v0.3")
    ap.add_argument("--data",default=str(ROOT/"data"/"sample_manual.json")); ap.add_argument("--output",default=str(ROOT/"output"/"zaitu_master_manual.docx")); ap.add_argument("--style",choices=sorted(STYLES),default="clean_blue"); ap.add_argument("--all-styles",action="store_true")
    a=ap.parse_args()
    if a.all_styles:
        for s in STYLES:
            render(a.data,Path(a.output).with_name(f"zaitu_master_{s}.docx"),s)
    else: render(a.data,a.output,a.style)

if __name__=="__main__": main()
